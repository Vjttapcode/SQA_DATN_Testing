# -*- coding: utf-8 -*-
"""Xuất 14_System-Test - Module lập lịch.csv -> .xlsx (ghi đè).

Tô nền đỏ nhạt các ô đã đổi so với bản gốc:
- Nếu có file baseline CSV (cùng cấu trúc): so sánh từng ô với baseline.
- Nếu không: suy ra ô đã qua sanitize black-box (đảo BLACKBOX_REPLACEMENTS + dòng thống kê QLND).

Đặt file gốc (trước khi sửa) tại:
  14_System-Test - Module lập lịch.baseline.csv
hoặc truyền đường dẫn: python export_system_test_xlsx.py path/to/baseline.csv
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from blackbox_rewrite_system_test import invert_sanitize_blackbox

ROOT = Path(__file__).parent
CSV_PATH = ROOT / "14_System-Test - Module lập lịch.csv"
XLSX_PATH = ROOT / "14_System-Test - Module lập lịch.xlsx"
BASELINE_DEFAULT = ROOT / "14_System-Test - Module lập lịch.baseline.csv"

RED_FILL = PatternFill(fill_type="solid", start_color="FFFFC7CE", end_color="FFFFC7CE")


def invert_stats_title(s: str) -> str:
    return (s or "").replace("Module Lập lịch (TKB)", "Module QLND")


def approx_inverted_cell(i: int, j: int, val: str) -> str:
    s = invert_sanitize_blackbox(val or "")
    if i == 0 and j == 0:
        s = invert_stats_title(s)
    return s


def change_mask_invert(df: pd.DataFrame) -> list[list[bool]]:
    n, m = df.shape
    mask: list[list[bool]] = []
    for i in range(n):
        row_m: list[bool] = []
        for j in range(m):
            val = str(df.iat[i, j])
            inv = approx_inverted_cell(i, j, val)
            row_m.append(inv != val)
        mask.append(row_m)
    return mask


def change_mask_baseline(df: pd.DataFrame, base: pd.DataFrame) -> list[list[bool]]:
    if df.shape != base.shape or list(df.columns) != list(base.columns):
        raise ValueError("baseline shape/columns must match current CSV")
    return (df != base).values.tolist()


def apply_red_highlights(xlsx_path: Path, mask: list[list[bool]]) -> None:
    wb = load_workbook(xlsx_path)
    ws = wb.active
    for i, row in enumerate(mask):
        for j, changed in enumerate(row):
            if changed:
                ws.cell(row=i + 2, column=j + 1).fill = RED_FILL
    wb.save(xlsx_path)


def main() -> None:
    if not CSV_PATH.is_file():
        raise SystemExit(f"Khong thay: {CSV_PATH}")

    baseline_path = Path(sys.argv[1]) if len(sys.argv) > 1 else BASELINE_DEFAULT
    df = pd.read_csv(CSV_PATH, encoding="utf-8-sig", dtype=str, keep_default_na=False)

    if baseline_path.is_file():
        try:
            base = pd.read_csv(baseline_path, encoding="utf-8-sig", dtype=str, keep_default_na=False)
            mask = change_mask_baseline(df, base)
            mode = "baseline"
        except ValueError:
            mask = change_mask_invert(df)
            mode = "invert (baseline mismatch)"
    else:
        mask = change_mask_invert(df)
        mode = "invert"

    df.to_excel(XLSX_PATH, index=False, engine="openpyxl")
    apply_red_highlights(XLSX_PATH, mask)

    n_red = sum(sum(r) for r in mask)
    print("OK rows", len(df), "cols", len(df.columns), "mode", mode, "cells_red", n_red)


if __name__ == "__main__":
    main()
