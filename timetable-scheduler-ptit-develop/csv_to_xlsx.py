#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CSV to XLSX Converter
Converts a CSV file to Excel XLSX format using openpyxl
"""

import csv
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("Error: openpyxl library is required.")
    print("Install it with: pip install openpyxl")
    sys.exit(1)


def convert_csv_to_xlsx(csv_file_path, xlsx_file_path, delimiter=','):
    """
    Convert CSV file to XLSX format

    Args:
        csv_file_path: Path to input CSV file
        xlsx_file_path: Path to output XLSX file
        delimiter: CSV delimiter (default comma)
    """
    import sys
    # Set UTF-8 encoding for console output on Windows
    if sys.platform == 'win32':
        import codecs
        sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'ignore')

    print(f"Converting: {csv_file_path}")
    print(f"Output: {xlsx_file_path}")

    # Create workbook and worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Read CSV and write to Excel
    with open(csv_file_path, 'r', encoding='utf-8-sig') as csvfile:
        reader = csv.reader(csvfile, delimiter=delimiter)
        for row_idx, row in enumerate(reader, start=1):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value.strip() if value else "")

                # Apply styles to header row
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                    cell.border = thin_border
                else:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Set row height for header
    ws.row_dimensions[1].height = 20

    # Freeze the header row
    ws.freeze_panes = "A2"

    # Save workbook
    wb.save(xlsx_file_path)
    print(f"\n✓ Conversion successful!")
    print(f"File saved: {xlsx_file_path}")


if __name__ == "__main__":
    # Default file names
    csv_file = "Unit Test - Lập lịch.csv"
    xlsx_file = "Unit Test - Lập lịch.xlsx"

    # Allow command line arguments
    if len(sys.argv) > 1:
        csv_file = sys.argv[1]
    if len(sys.argv) > 2:
        xlsx_file = sys.argv[2]

    convert_csv_to_xlsx(csv_file, xlsx_file)
